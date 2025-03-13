# coding:utf-8
import asyncio
import os

import matplotlib
import numpy as np
import scipy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm
from matplotlib import pyplot as plt
from openpyxl import Workbook

from compiled_functions import aot_calculations
from src.common.PermutationView import PermutationView
from src.common.TypeOfBasement import TypeOfBasement
from src.common.constants import Uz_a_0_16_x, Uz_a_0_16_z, Uz_a_0_25_x, Uz_a_0_25_z
from src.submodules.databasetoolkit.isolated import load_pressure_coefficients, find_experiment_by_model_name, \
    load_positions, load_face_number
from src.submodules.plot.plotBuilding import PlotBuilding
from src.submodules.plot.utils import scaling_data
from src.submodules.report_tools.reportFolder import ReportFolder
from src.submodules.report_tools.utils import create_directory_to_report
from src.submodules.report_tools.wordBuilder import WordBuilder
from src.submodules.utils.angle import get_angle_border, get_base_angle, changer_sequence_coefficients
from src.submodules.utils.data_features import lambdas, calculated, warranty_plus, warranty_minus, polar_lambdas
from src.submodules.utils.permutations import get_view_permutation_data, get_sequence_permutation_data
from src.submodules.utils.scaling import get_model_and_scale_factors
from src.submodules.utils.speed_sp import speed_sp_region
from src.submodules.utils.utils import converter_coordinates, get_size_and_count_sensors, tpu_size_to_real
from src.ui.common.ChartMode import ChartMode
from src.ui.common.ChartType import ChartType
from src.ui.common.CoordinateSystem import CoordinateSystem
from src.ui.common.IsofieldsType import IsofieldsType
from src.ui.view.interface import Interface


class IsolatedHighRiseInterface(Interface):
    """Isolated High Rise Interface

    constants :

    sample_frequency 32768
    sample_period 32.768

    """

    def __init__(
            self,
            parent=None,
            engine=None
    ):

        super().__init__(parent=parent, engine=engine)
        self.setObjectName('IsolatedHighRiseInterface')

    def get_pressure_coefficients_for_the_sensor(
            self,
            model_id,
            alpha,
            angle,
            model_name: str,
            sensor_id
    ):
        pressure_coefficients = self.get_pressure_coefficients(model_id, alpha, angle, str(model_name))
        return pressure_coefficients[:, sensor_id]

    def get_pressure_coefficients(
            self,
            model_id,
            alpha,
            angle,
            model_name: str
    ):
        model_name_base = model_name
        turn_flag = False
        if model_name[0] == model_name[1]:
            angle_border = 45
            type_base = TypeOfBasement.SQUARE

        else:
            angle_border = 90
            type_base = TypeOfBasement.RECTANGLE

        if model_name[1] in ['2', '3']:
            model_name_base = model_name[1] + model_name[0] + model_name[2]
            angle = str((int(angle) + 270) % 360)
            turn_flag = True

        # Поворот данных для отображения углов, выходящих за границы имеющихся
        if int(angle) > angle_border:
            permutation_view = get_view_permutation_data(type_base, int(angle))  # вид последовательности данных
            base_angle = get_base_angle(int(angle), permutation_view, type_base)
            sequence_permutation = get_sequence_permutation_data(type_base, permutation_view, int(angle))

            pressure_coefficients = asyncio.run(load_pressure_coefficients(model_id, alpha, self.engine,
                                                                           angle=base_angle))[base_angle]

            pressure_coefficients = changer_sequence_coefficients(pressure_coefficients, permutation_view,
                                                                  model_name, sequence_permutation)
        else:
            pressure_coefficients = asyncio.run(load_pressure_coefficients(model_id, alpha, self.engine,
                                                                           angle=angle))[angle]

        # Поворот модели
        if turn_flag:
            if int(angle) % 90 != 0:
                model_name_base = model_name
            pressure_coefficients = changer_sequence_coefficients(pressure_coefficients, PermutationView.FORWARD,
                                                                  model_name_base,
                                                                  (3, 0, 1, 2))
        return pressure_coefficients

    def get_model_id(
            self,
            model_name,
            alpha
    ):
        model_name_str = str(model_name)
        if model_name_str[1] in ['2', '3']:
            model_name = int(model_name_str[1] + model_name_str[0] + model_name_str[2])

        model_id = asyncio.run(find_experiment_by_model_name(model_name, alpha, self.engine)).model_id

        return model_id

    def get_coordinates(
            self,
            model_id,
            alpha
    ):
        coordinates = asyncio.run(load_positions(model_id, alpha, self.engine))

        return coordinates

    def get_face_number(
            self,
            model_id,
            alpha
    ):
        face_number = asyncio.run(load_face_number(model_id, alpha, self.engine))

        return face_number

    def create_report_sensor_statistics(
            self,
            coordinates,
            pressure_coefficients_storage: dict,
            count_sensors,
            face_number,
            angle_border,
            size_model_tpu,
            model_size,
            path_report

    ):
        breadth_tpu, depth_tpu, height_tpu = size_model_tpu
        breadth_real, depth_real, height_real = model_size
        x, z = coordinates

        x_new, y_new = converter_coordinates(x, breadth_tpu, depth_tpu, face_number, count_sensors, accuracy=5)

        headers = (
            'ДАТЧИК',
            'X(м)',
            'Y(м)',
            'Z(м)',
            ChartMode.MEAN,
            ChartMode.RMS,
            ChartMode.STD,
            ChartMode.MAX,
            ChartMode.MIN,
            ChartMode.CALCULATED,
            ChartMode.WARRANTY_PLUS,
            ChartMode.WARRANTY_MINUS
        )
        x_new = [tpu_size_to_real(x_new[sensor], breadth_real, breadth_tpu) for sensor in range(count_sensors)]
        y_new = [tpu_size_to_real(y_new[sensor], depth_real, depth_tpu) for sensor in range(count_sensors)]
        z_real = [tpu_size_to_real(z[sensor], height_real, height_tpu) for sensor in range(count_sensors)]

        wb = Workbook()

        # Удаляем лист по умолчанию, если он не нужен
        default_sheet = wb.active
        wb.remove(default_sheet)

        for angle in range(0, angle_border + 5, 5):
            sheet = wb.create_sheet(f"Угол {angle}")
            sheet.append(headers)
            pressure_coefficients = pressure_coefficients_storage[angle]

            statistics_of_angle = [[i for i in range(1, count_sensors + 1)]]

            statistics_of_angle.append(x_new)
            statistics_of_angle.append(y_new)
            statistics_of_angle.append(z_real)

            for k in lambdas:
                statistics_of_angle.append(lambdas[k](pressure_coefficients))

            statistics_of_angle.append(calculated(pressure_coefficients, axis=0))
            statistics_of_angle.append(warranty_plus(pressure_coefficients, axis=0))
            statistics_of_angle.append(warranty_minus(pressure_coefficients, axis=0))

            statistics_of_angle = np.array(statistics_of_angle).T

            for row in statistics_of_angle:
                sheet.append(row.tolist())

        wb.save(f'{os.path.join(path_report, ReportFolder.FILE_NAME_SENSOR_STATISTICS)}.xlsx')

    def draw_and_save_all_plots(
            self,
            pressure_coefficients_storage,
            alpha_str,
            coordinates,
            angle_border,
            path_report,
            model_size,
            model_size_str,
            model_name,
            size_model_tpu,
            count_sensors,
            wind_region,
    ):
        breadth_tpu, depth_tpu, height_tpu = size_model_tpu
        breadth, depth, height = model_size

        kz = height / size_model_tpu[2]

        match alpha_str:
            case 'A':
                Uz_a_x = Uz_a_0_16_x
                Uz_a_z = np.array(Uz_a_0_16_z)
            case 'C':
                Uz_a_x = Uz_a_0_25_x
                Uz_a_z = np.array(Uz_a_0_25_z)

        Uz_a_z_scaled = Uz_a_z * kz

        speed_tpu_function = scipy.interpolate.interp1d(Uz_a_z_scaled, Uz_a_x)
        speed_tpu = speed_tpu_function(height)

        speed_sp = speed_sp_region(height, alpha_str, wind_region)

        kv = speed_sp / speed_tpu

        # изополя в виде давления и коэффициентов
        for angle in range(0, angle_border + 5, 5):
            for parameter in (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS, ChartMode.STD):
                fig = PlotBuilding.isofields_coefficients(model_size,
                                                          model_name,
                                                          parameter,
                                                          pressure_coefficients_storage[angle],
                                                          coordinates)
                fig_name = f'{model_size_str} {alpha_str} {parameter} {angle}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.COEFFICIENT,
                                 parameter, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

                fig = PlotBuilding.isofields_coefficients(model_size,
                                                          model_name,
                                                          parameter,
                                                          pressure_coefficients_storage[angle],
                                                          coordinates,
                                                          alpha_str,
                                                          wind_region
                                                          )
                fig_name = f'{model_size_str} {alpha_str} {parameter} {angle}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.PRESSURE,
                                 parameter, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

        # изополя в виде мозаики

        for angle in range(0, angle_border + 5, 5):
            for parameter in (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS, ChartMode.STD):
                fig = PlotBuilding.pseudocolor_coefficients(model_size,
                                                            model_name,
                                                            parameter,
                                                            pressure_coefficients_storage[angle])
                fig_name = f'{model_size_str} {alpha_str} {parameter} {angle}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.DISCRETE_ISOFIELDS,
                                 parameter, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

        # огибающие

        for angle in range(0, angle_border + 5, 5):
            figs = PlotBuilding.envelopes(pressure_coefficients_storage[angle],
                                          (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS, ChartMode.STD))
            for i, fig in enumerate(figs):
                fig_name = f'{model_size_str} {alpha_str} {angle} {i}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.ENVELOPES, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

        data_to_plot = {}
        data_to_plot_polar = {}

        parameters = [
            ChartMode.MAX,
            ChartMode.MEAN,
            ChartMode.MIN,
            ChartMode.RMS,
            ChartMode.STD,
            ChartMode.CALCULATED,
            ChartMode.WARRANTY_PLUS,
            ChartMode.WARRANTY_MINUS,
        ]

        for v in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
            data_to_plot_polar[v] = {}
            for p in parameters:
                data_to_plot_polar[v][p] = []

        x = np.array(coordinates[0])
        y = np.array(coordinates[1])

        # получаем Cx Cy CMz для всех параметров

        for angle in range(0, angle_border + 5, 5):
            l_m = aot_calculations.calculate_projection_on_the_axis(breadth, depth, angle)
            l_tpu = aot_calculations.calculate_projection_on_the_axis(breadth_tpu, depth_tpu, angle)
            km = l_m / l_tpu
            kt = km / kv
            data_to_plot[angle] = {}
            cx, cy = aot_calculations.calculate_cx_cy(
                *count_sensors,
                *size_model_tpu,
                x,
                y,
                pressure_coefficients_storage[angle]
            )

            cmz = aot_calculations.calculate_cmz(
                *count_sensors,
                angle,
                *size_model_tpu,
                x,
                y,
                pressure_coefficients_storage[angle]
            )

            for p in parameters:
                data_to_plot_polar[ChartMode.CX][p].append(polar_lambdas[p](cx))
                data_to_plot_polar[ChartMode.CY][p].append(polar_lambdas[p](cy))
                data_to_plot_polar[ChartMode.CMZ][p].append(polar_lambdas[p](cmz))

            data_to_plot[angle][ChartMode.CX] = cx
            data_to_plot[angle][ChartMode.CY] = cy
            data_to_plot[angle][ChartMode.CMZ] = cmz
            # отрисовка CMz

            fig = PlotBuilding.summary_coefficients({ChartMode.CMZ: data_to_plot[angle][ChartMode.CMZ]},
                                                    kt)
            fig_name = f'{model_size_str} {alpha_str} {angle} {ChartMode.CMZ}.png'
            fig.set_size_inches(18.5, 10.5)
            fig.savefig(
                os.path.join(path_report, ChartType.SUMMARY_COEFFICIENTS,
                             CoordinateSystem.CARTESIAN, fig_name),
                dpi=200,
                bbox_inches='tight')
            plt.close(fig)
            # отрисовка Cx Cy

            fig = PlotBuilding.summary_coefficients({ChartMode.CX: data_to_plot[angle][ChartMode.CX],
                                                     ChartMode.CY: data_to_plot[angle][ChartMode.CY]}
                                                    , kt)
            fig_name = f'{model_size_str} {alpha_str} {angle} {ChartMode.CX} {ChartMode.CY}.png'
            fig.set_size_inches(18.5, 10.5)
            fig.savefig(
                os.path.join(path_report, ChartType.SUMMARY_COEFFICIENTS,
                             CoordinateSystem.CARTESIAN, fig_name),
                dpi=200,
                bbox_inches='tight')
            plt.close(fig)

        # Спектры
        for angle in range(0, angle_border + 5, 5):
            for i in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
                fig = PlotBuilding.welch_graph({i: data_to_plot[angle][i]}, height, speed_sp)
                fig_name = f'{model_size_str} {alpha_str} {angle}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.SPECTRUM, i, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

        del data_to_plot

        # Масштабируем данные для полярной системы координат

        for p in parameters:
            cx_scale, cy_scale = scaling_data(data_to_plot_polar[ChartMode.CX][p], data_to_plot_polar[ChartMode.CY][p],
                                              angle_border=angle_border)
            data_to_plot_polar[ChartMode.CX][p] = cx_scale
            data_to_plot_polar[ChartMode.CY][p] = cy_scale

            cmz_scale = scaling_data(data_to_plot_polar[ChartMode.CMZ][p], angle_border=angle_border)
            data_to_plot_polar[ChartMode.CMZ][p] = cmz_scale

        # Отрисовка в полярной системе координат

        for i in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
            for p in parameters:
                fig = PlotBuilding.polar_plot({i: {p: data_to_plot_polar[i][p]}})
                fig_name = f'{model_size_str} {alpha_str} {p}.png'
                fig.set_size_inches(18.5, 10.5)
                fig.savefig(
                    os.path.join(path_report, ChartType.SUMMARY_COEFFICIENTS,
                                 CoordinateSystem.POLAR, i, fig_name),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)

        del data_to_plot_polar

    def create_word_report(
            self,
            model_size,
            wind_region,
            alpha_str,
            path_report,
            report_name

    ):
        breadth, depth, height = model_size
        breadth = int(breadth) if breadth.is_integer() else f'{round(breadth, 2):.2f}'
        depth = int(depth) if depth.is_integer() else f'{round(depth, 2):.2f}'
        height = int(height) if height.is_integer() else f'{round(height, 2):.2f}'

        # Работа с word файлом
        doc = Document()
        style = doc.styles['Normal']
        style.font.size = Pt(14)
        style.font.name = 'Times New Roman'
        section = doc.sections[0]
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        # ширина A4 210 мм высота 297 мм
        fig_width = Mm(165)
        fig_height = Mm(297 / 2 - 55)

        # Шрифт заголовков разного уровня
        head_lvl2 = 16
        head_lvl3 = 16

        counter_plots = 1  # Счетчик графиков для нумерации
        counter_head_lvl1 = 1  # Счетчик заголовков

        WordBuilder.add_heading(doc,
                                head_name=f'Отчет по зданию {breadth}x{depth}x{height}',
                                font_size=24,
                                bold=True)

        for i in ('Параметры ветрового районирования:',
                  f'Ветровой район: {wind_region}',
                  f'Тип местности: {alpha_str}'
                  ):
            doc.add_paragraph().add_run(i)
        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}. Геометрические размеры здания',
                                bold=True)

        counter_head_lvl1 += 1

        table = [
            ['Геометрический размер', 'Значение, м'],
            ['Ширина:', breadth],
            ['Глубина:', depth],
            ['Высота:', height],
        ]
        WordBuilder.add_table(doc, table)

        # Создание содержания
        WordBuilder.add_heading(doc, head_name='Содержание', bold=True, page_break=True)

        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        fldChar = OxmlElement('w:fldChar')  # creates a new element
        fldChar.set(qn('w:fldCharType'), 'begin')  # sets attribute on element
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')  # sets attribute on element
        instrText.text = 'TOC \\o "1-3" \\h \\z \\u'  # change 1-3 depending on heading levels you need

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:t')
        fldChar3.text = "Right-click to update field."
        fldChar2.append(fldChar3)

        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')

        r_element = run._r
        r_element.append(fldChar)
        r_element.append(instrText)
        r_element.append(fldChar2)
        r_element.append(fldChar4)

        # огибающие

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.ENVELOPES}', page_break=True)
        counter_head_lvl1 += 1

        counter_plots = WordBuilder.fill_chapter_with_pictures(
            doc,
            folder_path=os.path.join(path_report, ChartType.ENVELOPES),
            counter_pictures=counter_plots,
            picture_height=fig_height
        )

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.ISOFIELDS}', page_break=True)

        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}. {ChartType.ISOFIELDS} {IsofieldsType.COEFFICIENT.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.COEFFICIENT)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.ISOFIELDS} {IsofieldsType.COEFFICIENT.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)

            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl2 += 1

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.ISOFIELDS} {IsofieldsType.PRESSURE.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.PRESSURE)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.ISOFIELDS} {IsofieldsType.PRESSURE.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)
            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl1 += 1
        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.SUMMARY_COEFFICIENTS}',
                                page_break=True)

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SUMMARY_COEFFICIENTS} {CoordinateSystem.CARTESIAN.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl2 += 1

        counter_plots = WordBuilder.fill_chapter_with_pictures(
            doc,
            folder_path=os.path.join(path_report, ChartType.SUMMARY_COEFFICIENTS, CoordinateSystem.CARTESIAN),
            counter_pictures=counter_plots,
            picture_height=fig_height - Mm(10)
        )

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SUMMARY_COEFFICIENTS} {CoordinateSystem.POLAR.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.SUMMARY_COEFFICIENTS, CoordinateSystem.POLAR)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.SUMMARY_COEFFICIENTS} {CoordinateSystem.POLAR.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)

            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl1 += 1
        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.SPECTRUM}', page_break=True)

        path_temp = os.path.join(path_report, ChartType.SPECTRUM)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SPECTRUM} {mode}',
                                    head_level=2,
                                    font_size=head_lvl2)

            counter_head_lvl2 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        doc.save(os.path.join(path_report, f'{report_name}.docx'))

    def create_report(
            self
    ):
        # need to switch backend to Agg to avoid memory leak
        matplotlib.use('Agg')

        model_size = self._get_model_size()
        alpha = self._get_alpha()
        alpha_str = self._get_alpha(string=True)
        wind_region = self._get_wind_region()

        model_name, _ = get_model_and_scale_factors(*model_size, alpha)
        model_id = self.get_model_id(model_name, alpha)
        coordinates = self.get_coordinates(model_id, alpha)
        face_number = self.get_face_number(model_id, alpha)

        model_size_str = " ".join(list(map(str, model_size)))
        report_name = f'{model_size_str} {alpha_str} {wind_region}'
        path_report = os.path.join(ReportFolder.WORD_REPORT, report_name)

        create_directory_to_report(report_name)
        angle_border = get_angle_border(str(model_name))
        size_model_tpu, count_sensors = get_size_and_count_sensors(len(coordinates[0]),
                                                                   model_name,
                                                                   )

        # Получаем коэффициенты сразу для всех углов
        pressure_coefficients_storage = {}
        for angle in range(0, angle_border + 5, 5):
            pressure_coefficients = self.get_pressure_coefficients(model_id, alpha, angle, str(model_name))

            pressure_coefficients_storage[angle] = pressure_coefficients

        self.create_report_sensor_statistics(coordinates, pressure_coefficients_storage, count_sensors[0],
                                             face_number,
                                             angle_border,
                                             size_model_tpu,
                                             model_size,
                                             path_report)
        self.draw_and_save_all_plots(pressure_coefficients_storage, alpha_str, coordinates, angle_border, path_report,
                                     model_size, model_size_str, model_name, size_model_tpu, count_sensors, wind_region)

        del pressure_coefficients_storage

        self.create_word_report(model_size, wind_region, alpha_str, path_report, report_name)
        # return default backend
        matplotlib.use('qtagg')
