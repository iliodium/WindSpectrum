# coding:utf-8
import os
from abc import abstractmethod
from concurrent.futures import ProcessPoolExecutor

import matplotlib
import numpy as np
from compiled_functions import aot_calculations
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import (Mm,
                         Pt,)
from matplotlib import pyplot as plt
from openpyxl import Workbook
from PySide6 import (QtCore,
                     QtGui,)
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (QHBoxLayout,
                               QStackedLayout,
                               QVBoxLayout,
                               QWidget,)
from qfluentwidgets import (ComboBox,
                            LineEdit,
                            PushButton,
                            StrongBodyLabel,
                            TitleLabel,)
from sqlalchemy import create_engine
from src.common.annotation import ModelSizeType
from src.common.constants import (alpha_standards,
                                  wind_regions,)
from src.submodules.plot.plotBuilding import PlotBuilding
from src.submodules.plot.utils import scaling_data
from src.submodules.report_tools.reportFolder import ReportFolder
from src.submodules.report_tools.utils import create_directory_to_report
from src.submodules.report_tools.wordBuilder import WordBuilder
from src.submodules.utils.data_features import (calculated,
                                                lambdas,
                                                polar_lambdas,
                                                warranty_minus,
                                                warranty_plus,)
from src.submodules.utils.scaling import calculate_kt
from src.submodules.utils.speed_sp import speed_sp_region
from src.submodules.utils.utils import (converter_coordinates,
                                        tpu_size_to_real,)
from src.ui.common.Buttons import Buttons
from src.ui.common.CartesianModelSummaryCoefficients import (
    CartesianModelSummaryCoefficients,)
from src.ui.common.ChartMode import ChartMode
from src.ui.common.ChartType import ChartType
from src.ui.common.CoordinateSystem import CoordinateSystem
from src.ui.common.IsofieldsType import IsofieldsType
from src.ui.common.StyleSheet import StyleSheet
from src.ui.components.MultiSelectComboBox import MultiSelectComboBox
from src.ui.view.widgets.MatplotlibWidget import MatplotlibWidget
from src.ui.view.widgets.SensorWidget import SensorWidget


class InterfaceBuildings(QWidget):
    """Interface"""

    SAMPLE_PERIOD = None
    SAMPLE_FREQUENCY = None
    NUMBER_OF_TIME_COUNTS = None

    REPORT_FOLDER_NAME = None
    # config
    DB_URL_LOCAL = None
    DB_URL_SERVER = None
    MAX_WORKERS = None

    def __init__(
            self,
            parent=None,
            config=None
    ):
        for key, value in config.items():
            setattr(self, key, value)
        super().__init__(parent=parent)
        self.engine = create_engine(self.DB_URL_LOCAL)
        self.view = self
        self.plotFlag = False

        # Set style
        StyleSheet.MAIN_INTERFACE.apply(self)

        self.hBoxLayoutMain = QHBoxLayout(self)

        # Initialization left menu
        self._init_general_information()

        self.StackedLayoutMainMenu = QStackedLayout()
        self.hBoxLayoutMain.addLayout(self.StackedLayoutMainMenu)

        # Initialization chart menu
        WidgetChartMenu = QWidget()
        self.vBoxLayoutPlot = QVBoxLayout(WidgetChartMenu)
        self.vBoxLayoutPlot.setAlignment(Qt.AlignTop | Qt.AlignLeft)

        self.StackedLayoutMainMenu.addWidget(WidgetChartMenu)

        self._init_chart_menu()

        WidgetSensorsOverview = QWidget()
        self.vBoxLayoutSensorsOverview = QVBoxLayout(WidgetSensorsOverview)
        self.vBoxLayoutSensorsOverview.setAlignment(Qt.AlignTop | Qt.AlignLeft)

        self.StackedLayoutMainMenu.addWidget(WidgetSensorsOverview)

        self._init_sensors_overview()

        # An array to store references to objects (like envelopes)
        # If you do not store the objects, they will be deleted by garbage collector
        self.plots = []

        self.fig_spectrum_sensors_overview = None
        self.fig_summary_coefficients_sensors_overview = None


    def _draw_sensors_overview(
            self,
            size_model_tpu,
            coordinates
    ):

        breadth_tpu, depth_tpu, height_tpu = size_model_tpu

        self.SensorWidget.model_size = size_model_tpu

        length = 2 * (breadth_tpu + depth_tpu)

        lines_pos = [i / length for i in (breadth_tpu, breadth_tpu + depth_tpu, 2 * breadth_tpu + depth_tpu)]

        self.SensorWidget.lines_pos = lines_pos
        self.SensorWidget.update_lines()

        length_x = 2 * (breadth_tpu + depth_tpu)
        length_y = height_tpu

        x = [i / length_x for i in coordinates[0]]
        y = [i / length_y for i in coordinates[1]]

        for b in self.SensorWidget.points:
            self.SensorWidget.scene.removeItem(b)
            b.deleteLater()  # Уничтожаем объект

        self.SensorWidget.points = []
        buttons_pos = [(i, 1 - j) for i, j in zip(x, y)]
        self.SensorWidget.points_pos = buttons_pos

        self.SensorWidget.add_points()

    def _switch_stacked_layout_sensors_overview(
            self
    ):
        current_index = self.StackedLayoutMainMenu.currentIndex()
        new_index = 1 if current_index == 0 else 0
        if new_index == 1:
            self.PushButtonSensorsOverview.setText(Buttons.PLOTS)
            coordinates = self._get_coordinates()
            size_model_tpu, count_sensors = self._get_size_and_count_sensors(len(coordinates[0]))
            if size_model_tpu != self.SensorWidget.model_size:
                self._draw_sensors_overview(size_model_tpu, coordinates)

            else:
                pass
        else:
            self.PushButtonSensorsOverview.setText(Buttons.SENSORS)

        self.StackedLayoutMainMenu.setCurrentIndex(new_index)

    @abstractmethod
    def _get_size_tpu(
            self
    ):
        pass

    def sensors_overview_button_action(self, sensor_id):
        model_size = self._get_model_size()

        pressure_coefficients = self._get_pressure_coefficients_for_the_sensor(sensor_id)
        type_plot = ChartType(self.ComboBoxSensorsOverview.currentText())

        angle = self._get_angle()
        alpha_str = self._get_alpha(area_type=True)
        wind_region = self._get_wind_region()
        size_tpu = self._get_size_tpu()

        match type_plot:
            case ChartType.AERODYNAMIC_COEFFICIENTS:
                kt = calculate_kt(model_size, size_tpu, alpha_str, wind_region, angle)

                if self.fig_summary_coefficients_sensors_overview is None:
                    fig = PlotBuilding.sensor_signal({f"Датчик {sensor_id + 1}": pressure_coefficients},
                                                     kt=kt,
                                                     sample_period=self.SAMPLE_PERIOD,
                                                     number_of_time_counts=self.NUMBER_OF_TIME_COUNTS)
                else:
                    data_to_plot = {f"Датчик {sensor_id + 1}": pressure_coefficients}
                    # Перенос данных из первой фигуры
                    for ax in self.fig_summary_coefficients_sensors_overview.get_axes():
                        for line in ax.get_lines():
                            data_to_plot[line.get_label()] = line.get_ydata()
                    self.windows_plot_sensors_overview_summary_coefficients.close()

                    fig = PlotBuilding.sensor_signal(data_to_plot,
                                                     kt=kt,
                                                     sample_period=self.SAMPLE_PERIOD,
                                                     number_of_time_counts=self.NUMBER_OF_TIME_COUNTS)

                self.fig_summary_coefficients_sensors_overview = fig
                self.open_plot_in_new_window_sensors_overview_summary_coefficients(fig,
                                                                                   ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS)
            case ChartType.SPECTRUM:
                height = model_size[2]
                z_coordinate = self._get_coordinates()[1][sensor_id]

                height_sensor = tpu_size_to_real(z_coordinate, height, size_tpu[2])
                speed_sp = speed_sp_region(height_sensor, alpha_str, wind_region)

                if self.fig_spectrum_sensors_overview is None:
                    fig = PlotBuilding.welch_graph({f"Датчик {sensor_id + 1}": pressure_coefficients},
                                                   height_sensor, speed_sp,
                                                   self.SAMPLE_FREQUENCY, self.NUMBER_OF_TIME_COUNTS)
                else:
                    data_to_plot = {f"Датчик {sensor_id + 1}": pressure_coefficients}
                    # Перенос данных из первой фигуры
                    for ax in self.fig_spectrum_sensors_overview.get_axes():
                        for line in ax.get_lines():
                            label = line.get_label()
                            last_space_index = label.rfind(" ")
                            sensor_id = int(label[last_space_index + 1:]) - 1
                            data_to_plot[label] = self._get_pressure_coefficients_for_the_sensor(sensor_id)

                    self.windows_plot_sensors_overview_summary_coefficients.close()

                    fig = PlotBuilding.welch_graph(data_to_plot, height_sensor, speed_sp,
                                                   self.SAMPLE_FREQUENCY, self.NUMBER_OF_TIME_COUNTS)

                self.fig_spectrum_sensors_overview = fig
                self.open_plot_in_new_window_sensors_overview_summary_coefficients(fig, ChartType.SPECTRUM)

    def _init_general_information(
            self
    ):
        self.generalInformationContainer = QWidget()
        self.generalInformationContainer.setFixedWidth(275)
        self.generalInformationContainer.setFixedHeight(300)

        self.vBoxLayoutGenInf = QVBoxLayout(self.generalInformationContainer)

        # Add label to grid layout
        self.vBoxLayoutGenInf.addWidget(TitleLabel(Buttons.GENERAL_INFORMATION))

        # Wind regions
        # Create horizontal box layout
        self.hBoxLayoutWindRegions = QHBoxLayout()
        # Add label to horizontal box layout
        self.hBoxLayoutWindRegions.addWidget(StrongBodyLabel(Buttons.WIND_REGION))
        # Create combo box
        self.ComboBoxWindRegions = ComboBox()
        # Fill the combo box
        self.ComboBoxWindRegions.addItems([
            self.tr(i) for i in [*wind_regions]
        ])
        # set fixed width of combobox
        self.ComboBoxWindRegions.setFixedWidth(75)
        # Add combo box to horizontal box layout
        self.hBoxLayoutWindRegions.addWidget(self.ComboBoxWindRegions)
        self.vBoxLayoutGenInf.addLayout(self.hBoxLayoutWindRegions)

        # Type of area
        self.hBoxLayoutTypeOfArea = QHBoxLayout(self.view)
        self.hBoxLayoutTypeOfArea.addWidget(StrongBodyLabel(Buttons.TYPE_OF_AREA))
        self.ComboBoxTypeOfArea = ComboBox()
        self.ComboBoxTypeOfArea.addItems([
            self.tr(i) for i in [*alpha_standards]
        ])
        self.ComboBoxTypeOfArea.setFixedWidth(75)
        self.hBoxLayoutTypeOfArea.addWidget(self.ComboBoxTypeOfArea)
        self.vBoxLayoutGenInf.addLayout(self.hBoxLayoutTypeOfArea)

        # Wind angle
        self.hBoxLayoutWindAngle = QHBoxLayout(self.view)
        self.hBoxLayoutWindAngle.addWidget(StrongBodyLabel(Buttons.WIND_ANGLE))
        # Create text input widget
        self.lineEditWindAngle = LineEdit()
        # Set default text
        self.lineEditWindAngle.setText(self.tr('0'))
        # Set clear button
        self.lineEditWindAngle.setClearButtonEnabled(True)
        self.lineEditWindAngle.setFixedWidth(75)
        # Add text input widget to horizontal box layout
        self.hBoxLayoutWindAngle.addWidget(self.lineEditWindAngle)
        self.vBoxLayoutGenInf.addLayout(self.hBoxLayoutWindAngle)

        # Building size
        self.hBoxLayoutBuildingSizeInterfering = QHBoxLayout(self.view)
        self.hBoxLayoutBuildingSizeInterfering.addWidget(StrongBodyLabel(Buttons.BUILDING_SIZE))
        self.lineEditBuildingSize = LineEdit()
        self.lineEditBuildingSize.setText(self.tr('10 10 20'))
        self.lineEditBuildingSize.setClearButtonEnabled(True)
        self.lineEditBuildingSize.setFixedWidth(125)
        self.hBoxLayoutBuildingSizeInterfering.addWidget(self.lineEditBuildingSize)
        self.vBoxLayoutGenInf.addLayout(self.hBoxLayoutBuildingSizeInterfering)

        PushButtonReport = PushButton(Buttons.REPORT)
        PushButtonReport.clicked.connect(self.create_report)
        self.vBoxLayoutGenInf.addWidget(PushButtonReport)

        self.PushButtonSensorsOverview = PushButton(Buttons.SENSORS)
        self.PushButtonSensorsOverview.clicked.connect(self._switch_stacked_layout_sensors_overview)
        self.vBoxLayoutGenInf.addWidget(self.PushButtonSensorsOverview)

        self.PushButtonInterferingInformation = PushButton(Buttons.FEA)
        # self.PushButtonFiniteElementMethod.clicked.connect(self._switch_stacked_layout_sensors_overview)
        self.vBoxLayoutGenInf.addWidget(self.PushButtonInterferingInformation)

        self.hBoxLayoutMain.addWidget(self.generalInformationContainer)

    def _init_chart_menu(
            self
    ):
        container = QWidget()
        container.setFixedWidth(775)
        container.setFixedHeight(100)

        # Chart menu
        hBoxLayoutChartMenu = QHBoxLayout(container)

        self.ComboBoxChartMenu = ComboBox()
        self.ComboBoxChartMenu.addItems([
            self.tr(i) for i in ChartType if i != ChartType.AERODYNAMIC_COEFFICIENTS
        ])
        self.ComboBoxChartMenu.currentTextChanged.connect(self._switch_stacked_layout_type_chart)
        self.ComboBoxChartMenu.setFixedWidth(225)
        hBoxLayoutChartMenu.addWidget(self.ComboBoxChartMenu)

        self.StackedLayoutTypeChart = QStackedLayout()

        self._init_chart_isofields()
        self._init_chart_envelopes()
        self._init_chart_summary_coefficients()
        self._init_chart_spectrum()
        self._init_chart_pseudocolor_coefficients()

        hBoxLayoutChartMenu.addLayout(self.StackedLayoutTypeChart)

        PushButtonCreatePlot = PushButton(Buttons.BUILD_PLOT)
        PushButtonCreatePlot.clicked.connect(self.create_plot)
        PushButtonCreatePlot.setFixedWidth(100)

        hBoxLayoutChartMenu.addWidget(PushButtonCreatePlot)
        self.vBoxLayoutPlot.addWidget(container)

    def _init_sensors_overview(self):
        container = QWidget()
        container.setFixedWidth(775)
        container.setFixedHeight(100)

        # Chart menu
        hBoxLayoutChartMenu = QHBoxLayout(container)
        hBoxLayoutChartMenu.setAlignment(Qt.AlignLeft)
        self.ComboBoxSensorsOverview = ComboBox()
        self.ComboBoxSensorsOverview.addItems([
            self.tr(i) for i in (ChartType.AERODYNAMIC_COEFFICIENTS, ChartType.SPECTRUM)
        ])
        self.ComboBoxSensorsOverview.setFixedWidth(300)
        hBoxLayoutChartMenu.addWidget(self.ComboBoxSensorsOverview)
        self.vBoxLayoutSensorsOverview.addWidget(container)

        self.SensorWidget = SensorWidget(self.sensors_overview_button_action)
        self.vBoxLayoutSensorsOverview.addWidget(self.SensorWidget.view)

    def create_plot(
            self
    ):
        match self.StackedLayoutTypeChart.currentIndex():
            case 0:
                self.plot_isofields()
            case 1:
                self.plot_envelopes()
            case 2:
                self.plot_summary_coefficients()
            case 3:
                self.plot_welch_graph()
            case 4:
                self.plot_pseudocolor_coefficients()

    def _init_chart_isofields(
            self
    ):
        self.WidgetIsofields = QWidget()
        self.hBoxLayoutIsofields = QHBoxLayout(self.WidgetIsofields)

        self.ComboBoxTypesIsofields = ComboBox()
        self.ComboBoxTypesIsofields.addItems([
            self.tr(i) for i in (IsofieldsType.PRESSURE,
                                 IsofieldsType.COEFFICIENT,
                                 )
        ])
        self.ComboBoxTypesIsofields.setFixedWidth(140)
        self.hBoxLayoutIsofields.addWidget(self.ComboBoxTypesIsofields)

        self.isofieldsParameters = ComboBox()
        self.isofieldsParameters.addItems([
            self.tr(i) for i in (ChartMode.MAX,
                                 ChartMode.MEAN,
                                 ChartMode.MIN,
                                 ChartMode.RMS,
                                 ChartMode.STD,
                                 )])
        self.hBoxLayoutIsofields.addWidget(self.isofieldsParameters)

        self.StackedLayoutTypeChart.addWidget(self.WidgetIsofields)

    def _init_chart_envelopes(
            self
    ):
        self.WidgetEnvelopes = QWidget()
        self.hBoxLayoutEnvelopes = QHBoxLayout(self.WidgetEnvelopes)
        self.envelopesParameters = MultiSelectComboBox(placeholderText=Buttons.PARAMETERS)
        self.envelopesParameters.addItems([ChartMode.MAX,
                                           ChartMode.MEAN,
                                           ChartMode.MIN,
                                           ChartMode.RMS,
                                           ChartMode.STD,
                                           ])
        # self.envelopesParameters.checkAllItems()
        self.hBoxLayoutEnvelopes.addWidget(self.envelopesParameters)
        self.StackedLayoutTypeChart.addWidget(self.WidgetEnvelopes)

    def _init_chart_summary_coefficients(
            self
    ):
        self.WidgetSummaryCoefficients = QWidget()
        self.hBoxLayoutSummaryCoefficients = QHBoxLayout(self.WidgetSummaryCoefficients)

        # Combo box with coordinate system
        self.ComboBoxCoordinateSystemSummaryCoefficients = ComboBox()
        self.ComboBoxCoordinateSystemSummaryCoefficients.addItems([
            self.tr(i) for i in (CoordinateSystem.CARTESIAN,
                                 CoordinateSystem.POLAR,
                                 )
        ])
        self.ComboBoxCoordinateSystemSummaryCoefficients.currentTextChanged.connect(
            self._switch_stacked_layout_summary_coefficients)
        self.ComboBoxCoordinateSystemSummaryCoefficients.setFixedWidth(140)
        self.hBoxLayoutSummaryCoefficients.addWidget(self.ComboBoxCoordinateSystemSummaryCoefficients)
        # Stacked layout which varies depending on the coordinate system
        self.StackedLayoutSummaryCoefficients = QStackedLayout()
        self.hBoxLayoutSummaryCoefficients.addLayout(self.StackedLayoutSummaryCoefficients)
        # Cartesian system
        self.WidgetCartesianSummaryCoefficients = QWidget()
        self.hBoxLayoutCartesianCoordinateSystem = QHBoxLayout(self.WidgetCartesianSummaryCoefficients)
        self.cartesianModelSummaryCoefficients = ComboBox()
        self.cartesianModelSummaryCoefficients.addItems([
            self.tr(i) for i in (CartesianModelSummaryCoefficients.REAL,
                                 CartesianModelSummaryCoefficients.TPU,
                                 )
        ])

        self.cartesianParameters = MultiSelectComboBox(placeholderText=Buttons.PARAMETERS)
        self.cartesianParameters.addItems([ChartMode.CX,
                                           ChartMode.CY,
                                           ChartMode.CMZ,
                                           ])
        self.hBoxLayoutCartesianCoordinateSystem.addWidget(self.cartesianModelSummaryCoefficients)
        self.hBoxLayoutCartesianCoordinateSystem.addWidget(self.cartesianParameters)
        self.StackedLayoutSummaryCoefficients.addWidget(self.WidgetCartesianSummaryCoefficients)
        # Polar system
        self.WidgetPolarSummaryCoefficients = QWidget()
        self.hBoxLayoutPolarCoordinateSystem = QHBoxLayout(self.WidgetPolarSummaryCoefficients)
        self.polarView = MultiSelectComboBox(placeholderText=Buttons.VIEW)
        self.polarView.addItems([ChartMode.CX,
                                 ChartMode.CY,
                                 ChartMode.CMZ,
                                 ])
        self.polarParameters = MultiSelectComboBox(placeholderText=Buttons.PARAMETERS)
        self.polarParameters.addItems([ChartMode.MAX,
                                       ChartMode.MEAN,
                                       ChartMode.MIN,
                                       ChartMode.RMS,
                                       ChartMode.STD,
                                       ChartMode.CALCULATED,
                                       ChartMode.WARRANTY_PLUS,
                                       ChartMode.WARRANTY_MINUS,
                                       ])
        self.hBoxLayoutPolarCoordinateSystem.addWidget(self.polarView)
        self.hBoxLayoutPolarCoordinateSystem.addWidget(self.polarParameters)
        self.StackedLayoutSummaryCoefficients.addWidget(self.WidgetPolarSummaryCoefficients)
        # Add widget to main layout
        self.StackedLayoutTypeChart.addWidget(self.WidgetSummaryCoefficients)

    def _init_chart_spectrum(
            self
    ):
        self.WidgetSpectrum = QWidget()
        self.hBoxLayoutSpectrum = QHBoxLayout(self.WidgetSpectrum)
        self.spectrumParameters = MultiSelectComboBox(placeholderText=Buttons.PARAMETERS)
        self.spectrumParameters.addItems([ChartMode.CX,
                                          ChartMode.CY,
                                          ChartMode.CMZ,
                                          ])
        self.hBoxLayoutSpectrum.addWidget(self.spectrumParameters)
        self.StackedLayoutTypeChart.addWidget(self.WidgetSpectrum)

    def _init_chart_pseudocolor_coefficients(
            self
    ):
        self.WidgetDiscreteIsofields = QWidget()
        self.hBoxLayoutDiscreteIsofields = QHBoxLayout(self.WidgetDiscreteIsofields)
        self.discreteIsofieldsParameters = ComboBox()
        self.discreteIsofieldsParameters.addItems([
            self.tr(i) for i in (ChartMode.MAX,
                                 ChartMode.MEAN,
                                 ChartMode.MIN,
                                 ChartMode.RMS,
                                 ChartMode.STD)
        ])
        self.hBoxLayoutDiscreteIsofields.addWidget(self.discreteIsofieldsParameters)
        self.StackedLayoutTypeChart.addWidget(self.WidgetDiscreteIsofields)

    def _switch_stacked_layout_type_chart(
            self,
            chart_type
    ):
        match chart_type:
            case ChartType.ISOFIELDS:
                self.StackedLayoutTypeChart.setCurrentIndex(0)
            case ChartType.ENVELOPES:
                self.StackedLayoutTypeChart.setCurrentIndex(1)
            case ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS:
                self.StackedLayoutTypeChart.setCurrentIndex(2)
            case ChartType.SPECTRUM:
                self.StackedLayoutTypeChart.setCurrentIndex(3)
            case ChartType.DISCRETE_ISOFIELDS:
                self.StackedLayoutTypeChart.setCurrentIndex(4)

    def _switch_stacked_layout_summary_coefficients(
            self,
            system
    ):
        match system:
            case CoordinateSystem.CARTESIAN:
                self.StackedLayoutSummaryCoefficients.setCurrentIndex(0)
            case CoordinateSystem.POLAR:
                self.StackedLayoutSummaryCoefficients.setCurrentIndex(1)

    def _get_model_size(
            self
    ) -> ModelSizeType:
        return tuple(map(float, self.lineEditBuildingSize.text().replace(',', '.').split(' ')))

    def _get_alpha(
            self,
            area_type=False
    ):
        if area_type:
            return self.ComboBoxTypeOfArea.text()

        else:
            type_alpha = {
                'A': 4,
                'C': 6,
            }
            return type_alpha[self.ComboBoxTypeOfArea.text()]

    def _get_wind_region(
            self
    ):

        return self.ComboBoxWindRegions.text()

    def _get_angle(
            self
    ):
        """получаем угол и возвращаем ближайшее кратное число на 5"""
        angle = int(self.lineEditWindAngle.text()) % 360

        return 5 * round(angle / 5)

    def _icon(
            self,
            path
    ):
        """
        link to the original ->
        matplotlib.backends.backend_qt.NavigationToolbar2QT._icon
        """
        pm = QtGui.QPixmap(path)
        pm.setDevicePixelRatio(
            self.devicePixelRatioF() or 1)  # rarely, devicePixelRatioF=0
        if self.palette().color(self.backgroundRole()).value() < 128:
            icon_color = self.palette().color(self.foregroundRole())
            mask = pm.createMaskFromColor(
                QtGui.QColor('black'),
                QtCore.Qt.MaskMode.MaskOutColor)
            pm.fill(icon_color)
            pm.setMask(mask)
        return QtGui.QIcon(pm)

    def open_in_new_window(
            self
    ):
        window = QWidget()
        window.setWindowTitle(self.plotWidget.plotTitle)
        self.plots.append(window)
        layout = QVBoxLayout(window)

        layout.addWidget(self.plotWidget.toolbar)
        layout.addWidget(self.plotWidget)

        self.plotFlag = False

        window.show()

    def del_plot(self):
        self.vBoxLayoutPlot.removeWidget(self.containerPlot)
        self.plotWidget.setParent(None)  # Убираем из иерархии
        self.plotWidget.toolbar.deleteLater()  # Уничтожаем объект
        self.plotWidget.deleteLater()  # Уничтожаем объект
        self.containerPlot = None

    def add_plot_on_screen(
            self,
            fig,
            title='График',
            open_in_new_window_button=True
    ):
        if self.plotFlag:
            self.del_plot()

        self.containerPlot = QWidget()
        self.containerPlot.setStyleSheet("border: 0px;")
        self.vBoxLayoutPlot1 = QVBoxLayout(self.containerPlot)

        fig.patch.set_facecolor('#f7f9fc')

        # Create plot widget
        self.plotWidget = MatplotlibWidget(fig, title)

        if open_in_new_window_button:
            path = r"src\ui\resource\images\fl_icon.png"
            act = self.plotWidget.toolbar.addAction(self._icon(path), 'Открыть в новом окне', self.open_in_new_window)
            # if you set the value to -1, the button will be in the rightmost position
            self.plotWidget.toolbar.insertAction(self.plotWidget.toolbar.actions()[-2], act)

        self.vBoxLayoutPlot1.addWidget(self.plotWidget.toolbar)
        self.vBoxLayoutPlot1.addWidget(self.plotWidget)
        self.vBoxLayoutPlot1.setAlignment(Qt.AlignTop | Qt.AlignLeft)

        self.vBoxLayoutPlot.addWidget(self.containerPlot)

        self.plotFlag = True

    def plot_welch_graph(
            self
    ):

        parameters = [ChartMode(i) for i in self.spectrumParameters.getCurrentOptions()]
        if not parameters:
            return

        data_to_plot = {}

        coordinates = self._get_coordinates()
        angle = self._get_angle()
        pressure_coefficients = self._get_pressure_coefficients()

        size_tpu, count_sensors = self._get_size_and_count_sensors(len(coordinates[0]))

        model_size = self._get_model_size()
        height = model_size[2]

        if ChartMode.CX in parameters or ChartMode.CY in parameters:
            cx, cy = aot_calculations.calculate_cx_cy(
                *count_sensors,
                *size_tpu,
                np.array(coordinates[0]),
                np.array(coordinates[1]),
                pressure_coefficients
            )
            if ChartMode.CX in parameters:
                data_to_plot[ChartMode.CX] = cx
            if ChartMode.CY in parameters:
                data_to_plot[ChartMode.CY] = cy

        if ChartMode.CMZ in parameters:
            cmz = aot_calculations.calculate_cmz(
                *count_sensors,
                angle,
                *size_tpu,
                np.array(coordinates[0]),
                np.array(coordinates[1]),
                pressure_coefficients
            )
            data_to_plot[ChartMode.CMZ] = cmz

        alpha_str = self._get_alpha(area_type=True)
        wind_region = self._get_wind_region()
        speed_sp = speed_sp_region(height, alpha_str, wind_region)

        fig = PlotBuilding.welch_graph(data_to_plot, height, speed_sp, sample_frequency=self.SAMPLE_FREQUENCY,
                                       number_of_time_counts=self.NUMBER_OF_TIME_COUNTS)
        self.add_plot_on_screen(fig, ChartType.ISOFIELDS)

    def open_plot_in_new_window(
            self,
            fig,
            title='График'
    ):
        window = QWidget()
        window.setWindowTitle(title)

        self.plots.append(window)
        layout = QVBoxLayout(window)
        plotWidget = MatplotlibWidget(fig)

        layout.addWidget(plotWidget.toolbar)
        layout.addWidget(plotWidget)
        window.show()

    def open_plot_in_new_window_sensors_overview_summary_coefficients(
            self,
            fig,
            title='График'
    ):
        self.windows_plot_sensors_overview_summary_coefficients = QWidget()
        self.windows_plot_sensors_overview_summary_coefficients.setWindowTitle(title)

        def closeEvent(event):
            match self.windows_plot_sensors_overview_summary_coefficients.windowTitle():
                case ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS:
                    self.fig_summary_coefficients_sensors_overview = None
                case ChartType.SPECTRUM:
                    self.fig_spectrum_sensors_overview = None
            event.accept()  # Закрываем окно

        # Привязываем кастомный closeEvent к окну
        self.windows_plot_sensors_overview_summary_coefficients.closeEvent = closeEvent

        self.plots.append(self.windows_plot_sensors_overview_summary_coefficients)
        layout = QVBoxLayout(self.windows_plot_sensors_overview_summary_coefficients)
        plotWidget = MatplotlibWidget(fig)

        layout.addWidget(plotWidget.toolbar)
        layout.addWidget(plotWidget)
        self.windows_plot_sensors_overview_summary_coefficients.show()

    def open_plot_in_new_window_sensors_overview_spectrum(
            self,
            fig,
            title='График'
    ):
        self.windows_plot_sensors_overview_spectrum = QWidget()
        self.windows_plot_sensors_overview_spectrum.setWindowTitle(title)

        def closeEvent(event):
            match self.windows_plot_sensors_overview_spectrum.windowTitle():
                case ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS:
                    self.fig_summary_coefficients_sensors_overview = None
                case ChartType.SPECTRUM:
                    self.fig_spectrum_sensors_overview = None
            event.accept()  # Закрываем окно

        # Привязываем кастомный closeEvent к окну
        self.windows_plot_sensors_overview_spectrum.closeEvent = closeEvent

        self.plots.append(self.windows_plot_sensors_overview_spectrum)
        layout = QVBoxLayout(self.windows_plot_sensors_overview_spectrum)
        plotWidget = MatplotlibWidget(fig)

        layout.addWidget(plotWidget.toolbar)
        layout.addWidget(plotWidget)
        self.windows_plot_sensors_overview_spectrum.show()

    def _plot_envelopes(
            self,
            pressure_coefficients,
            mods
    ):
        figs = PlotBuilding.envelopes(pressure_coefficients, mods)

        return figs

    def plot_envelopes(
            self
    ):
        mods = [ChartMode(i) for i in self.envelopesParameters.getCurrentOptions()]
        pressure_coefficients = self._get_pressure_coefficients()

        figs = self._plot_envelopes(pressure_coefficients, mods)

        for fig in figs:
            self.open_plot_in_new_window(fig, ChartType.ENVELOPES)

    @abstractmethod
    def _get_pressure_coefficients(
            self,
            *args,
            **kwargs
    ):
        pass

    def _get_pressure_coefficients_for_the_sensor(
            self,
            sensor_id
    ):
        pressure_coefficients = self._get_pressure_coefficients()

        return pressure_coefficients[:, sensor_id]

    @abstractmethod
    def get_model_id(self,
                     model_name,
                     alpha
                     ):
        pass

    @abstractmethod
    def _get_coordinates(
            self,
            *args,
            **kwargs
    ):
        pass

    @abstractmethod
    def _get_face_number(
            self
    ) -> list[int]:
        pass

    @abstractmethod
    def _get_pressure_coefficients_for_definition_angle_future(
            self,
            *args,
            **kwargs
    ):
        pass

    @staticmethod
    @abstractmethod
    def _get_pressure_coefficients_for_definition_angle(
            engine,
            db_url_server,
            *args,
            **kwargs
    ):
        pass

    @abstractmethod
    def _get_size_and_count_sensors(
            self,
            *args,
            **kwargs
    ):
        pass

    @abstractmethod
    def _get_model_name(
            self
    ):
        pass

    def _plot_isofields(
            self,
            model_size,
            size,
            count_sensors,
            parameter,
            pressure_coefficients,
            coordinates
    ):
        match self.ComboBoxTypesIsofields.text():
            case IsofieldsType.PRESSURE:
                area_type = self._get_alpha(area_type=True)
                wind_region = self._get_wind_region()
                fig = PlotBuilding.isofields_coefficients(model_size,
                                                          size,
                                                          count_sensors,
                                                          parameter,
                                                          pressure_coefficients,
                                                          coordinates,
                                                          area_type,
                                                          wind_region
                                                          )
            case IsofieldsType.COEFFICIENT:
                fig = PlotBuilding.isofields_coefficients(model_size,
                                                          size,
                                                          count_sensors,
                                                          parameter,
                                                          pressure_coefficients,
                                                          coordinates
                                                          )

        return fig

    def plot_isofields(
            self
    ):
        pressure_coefficients = self._get_pressure_coefficients()
        coordinates = self._get_coordinates()
        size, count_sensors = self._get_size_and_count_sensors(pressure_coefficients.shape[1])
        model_size = self._get_model_size()
        parameter = ChartMode(self.isofieldsParameters.currentText())

        fig = self._plot_isofields(model_size, size, count_sensors, parameter, pressure_coefficients, coordinates)

        self.add_plot_on_screen(fig, ChartType.ISOFIELDS)

    def plot_summary_coefficients_cartesian(
            self,
            angle,
            pressure_coefficients_storage: dict,
            coordinates,
            size_tpu,
            count_sensors,
            model_size
    ):

        data_to_plot = {}

        parameters = [ChartMode(i) for i in self.cartesianParameters.getCurrentOptions()]
        pressure_coefficients = pressure_coefficients_storage[angle]

        if ChartMode.CX in parameters or ChartMode.CY in parameters:
            cx, cy = aot_calculations.calculate_cx_cy(
                *count_sensors,
                *size_tpu,
                np.array(coordinates[0]),
                np.array(coordinates[1]),
                pressure_coefficients
            )
            if ChartMode.CX in parameters:
                data_to_plot[ChartMode.CX] = cx
            if ChartMode.CY in parameters:
                data_to_plot[ChartMode.CY] = cy
        if ChartMode.CMZ in parameters:
            cmz = aot_calculations.calculate_cmz(
                *count_sensors,
                angle,
                *size_tpu,
                np.array(coordinates[0]),
                np.array(coordinates[1]),
                pressure_coefficients
            )

            data_to_plot[ChartMode.CMZ] = cmz

        match CartesianModelSummaryCoefficients(self.cartesianModelSummaryCoefficients.currentText()):
            case CartesianModelSummaryCoefficients.TPU:
                kt = 1
            case CartesianModelSummaryCoefficients.REAL:
                alpha_str = self._get_alpha(area_type=True)
                wind_region = self._get_wind_region()

                kt = calculate_kt(model_size, size_tpu, alpha_str, wind_region, angle)

        fig = PlotBuilding.summary_coefficients(data_to_plot,
                                                kt,
                                                sample_period=self.SAMPLE_PERIOD,
                                                number_of_time_counts=self.NUMBER_OF_TIME_COUNTS)

        self.add_plot_on_screen(fig, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS)

    def plot_summary_coefficients_polar(
            self,
            angle_border,
            pressure_coefficients_storage: dict,
            coordinates,
            size_tpu,
            count_sensors
    ):

        data_to_plot = {}

        views = [ChartMode(i) for i in self.polarView.getCurrentOptions()]
        parameters = [ChartMode(i) for i in self.polarParameters.getCurrentOptions()]

        x = np.array(coordinates[0])
        y = np.array(coordinates[1])

        cx_flag = ChartMode.CX in views
        cy_flag = ChartMode.CY in views
        cmz_flag = ChartMode.CMZ in views

        if cx_flag or cy_flag:
            for v in [ChartMode.CX, ChartMode.CY]:
                data_to_plot[v] = {}
                for p in parameters:
                    data_to_plot[v][p] = []

        if cmz_flag:
            data_to_plot[ChartMode.CMZ] = {}
            for p in parameters:
                data_to_plot[ChartMode.CMZ][p] = []

        for angle in range(0, angle_border + 5, 5):
            pressure_coefficients = pressure_coefficients_storage[angle]

            if cx_flag or cy_flag:
                cx, cy = aot_calculations.calculate_cx_cy(
                    *count_sensors,
                    *size_tpu,
                    x,
                    y,
                    pressure_coefficients
                )
                for p in parameters:
                    data_to_plot[ChartMode.CX][p].append(polar_lambdas[p](cx))
                    data_to_plot[ChartMode.CY][p].append(polar_lambdas[p](cy))

            if cmz_flag:
                cmz = aot_calculations.calculate_cmz(
                    *count_sensors,
                    angle,
                    *size_tpu,
                    x,
                    y,
                    pressure_coefficients
                )
                for p in parameters:
                    data_to_plot[ChartMode.CMZ][p].append(polar_lambdas[p](cmz))

        if len(pressure_coefficients_storage) != 72:
            if cx_flag or cy_flag:
                for p in parameters:
                    cx_scale, cy_scale = scaling_data(data_to_plot[ChartMode.CX][p], data_to_plot[ChartMode.CY][p],
                                                      angle_border=angle_border)
                    data_to_plot[ChartMode.CX][p] = cx_scale
                    data_to_plot[ChartMode.CY][p] = cy_scale

            if cmz_flag:
                for p in parameters:
                    cmz_scale = scaling_data(data_to_plot[ChartMode.CMZ][p], angle_border=angle_border)
                    data_to_plot[ChartMode.CMZ][p] = cmz_scale
        else:
            if cx_flag:
                for p in parameters:
                    data_to_plot[ChartMode.CX][p] = np.append(data_to_plot[ChartMode.CX][p],
                                                              data_to_plot[ChartMode.CX][p][0])

            if cy_flag:
                for p in parameters:
                    data_to_plot[ChartMode.CY][p] = np.append(data_to_plot[ChartMode.CY][p],
                                                              data_to_plot[ChartMode.CY][p][0])

            if cmz_flag:
                for p in parameters:
                    data_to_plot[ChartMode.CMZ][p] = np.append(data_to_plot[ChartMode.CMZ][p],
                                                               data_to_plot[ChartMode.CMZ][p][0])

        if not cx_flag and ChartMode.CX in data_to_plot:
            del data_to_plot[ChartMode.CX]

        if not cy_flag and ChartMode.CY in data_to_plot:
            del data_to_plot[ChartMode.CY]

        if not cmz_flag and ChartMode.CMZ in data_to_plot:
            del data_to_plot[ChartMode.CMZ]

        fig = PlotBuilding.polar_plot(data_to_plot)

        self.add_plot_on_screen(fig, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS)

    def plot_summary_coefficients(
            self
    ):
        if not ([ChartMode(i) for i in self.cartesianParameters.getCurrentOptions()] or
                ([ChartMode(i) for i in self.polarView.getCurrentOptions()] and
                 [ChartMode(i) for i in self.polarView.getCurrentOptions()])):
            return

        type_plot = CoordinateSystem(self.ComboBoxCoordinateSystemSummaryCoefficients.currentText())

        coordinates = self._get_coordinates()

        size_tpu, count_sensors = self._get_size_and_count_sensors(len(coordinates[0]))

        pressure_coefficients_storage = {}

        match type_plot:
            case CoordinateSystem.CARTESIAN:
                model_size = self._get_model_size()
                angle = self._get_angle()

                pressure_coefficients = self._get_pressure_coefficients()
                pressure_coefficients_storage[angle] = pressure_coefficients

                self.plot_summary_coefficients_cartesian(angle, pressure_coefficients_storage, coordinates, size_tpu,
                                                         count_sensors, model_size)

            case CoordinateSystem.POLAR:
                pressure_coefficients_storage = self._get_pressure_coefficients_storage()
                angle_border = self._get_angle_border()
                self.plot_summary_coefficients_polar(angle_border, pressure_coefficients_storage, coordinates,
                                                     size_tpu, count_sensors)

    @abstractmethod
    def _get_pressure_coefficients_storage(
            self,
            *args,
            **kwargs
    ):
        pass

    @abstractmethod
    def _get_angle_border(
            self,
            *args,
            **kwargs
    ):
        pass

    def _plot_pseudocolor_coefficients(
            self,
            model_size,
            count_sensors,
            parameter,
            pressure_coefficients
    ):
        fig = PlotBuilding.pseudocolor_coefficients(model_size,
                                                    count_sensors,
                                                    parameter,
                                                    pressure_coefficients)

        return fig

    def plot_pseudocolor_coefficients(
            self
    ):
        pressure_coefficients = self._get_pressure_coefficients()
        size, count_sensors = self._get_size_and_count_sensors(pressure_coefficients.shape[1])
        model_size = self._get_model_size()
        parameter = ChartMode(self.discreteIsofieldsParameters.currentText())

        fig = self._plot_pseudocolor_coefficients(model_size, count_sensors, parameter, pressure_coefficients)

        self.add_plot_on_screen(fig, ChartType.DISCRETE_ISOFIELDS)

    @staticmethod
    def _draw_and_save_plot_future(
            func,
            path,
            fig_name,
            width,
            height,
            dpi,
            bbox_inches,
            *args
    ):
        figs = func(*args)
        if isinstance(figs, list):
            for i, fig in enumerate(figs):
                fig.set_size_inches(width, height)
                fig.savefig(
                    os.path.join(path, f'{fig_name} {i}.png'),
                    dpi=200,
                    bbox_inches='tight')
                plt.close(fig)
        else:
            figs.set_size_inches(width, height)
            figs.savefig(os.path.join(path, f'{fig_name}.png'), dpi=dpi, bbox_inches=bbox_inches)
            plt.close(figs)

    @staticmethod
    def _run_future(
            executor,
            *args
    ):
        future = executor.submit(*args)
        future.add_done_callback(
            lambda f: print(f"Ошибка в задаче: {f.exception()}") if f.exception() else None)

        return future

    def draw_and_save_all_plots(
            self,
            pressure_coefficients_storage,
            area_type,
            coordinates,
            angle_border,
            path_report,
            model_size,
            model_size_str,
            size_tpu,
            count_sensors,
            wind_region,
    ):
        # изополя в виде давления и коэффициентов
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for angle in range(0, angle_border + 5, 5):
                for parameter in (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS, ChartMode.STD):
                    fig_name = f'{model_size_str} {area_type} {parameter} {angle}'
                    path = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.COEFFICIENT,
                                        parameter)
                    args = (model_size,
                            size_tpu,
                            count_sensors,
                            parameter,
                            pressure_coefficients_storage[angle],
                            coordinates
                            )

                    self._run_future(executor, self._draw_and_save_plot_future,
                                     PlotBuilding.isofields_coefficients, path, fig_name, 18.5, 10.5, 200, 'tight',
                                     *args)

                    path = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.PRESSURE,
                                        parameter)
                    args = (model_size,
                            size_tpu,
                            count_sensors,
                            parameter,
                            pressure_coefficients_storage[angle],
                            coordinates,
                            area_type,
                            wind_region
                            )

                    self._run_future(executor, self._draw_and_save_plot_future,
                                     PlotBuilding.isofields_coefficients, path, fig_name, 18.5, 10.5, 200, 'tight',
                                     *args)

            executor.shutdown(wait=True)
        # изополя в виде мозаики
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for angle in range(0, angle_border + 5, 5):
                for parameter in (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS, ChartMode.STD):
                    fig_name = f'{model_size_str} {area_type} {parameter} {angle}'
                    path = os.path.join(path_report, ChartType.DISCRETE_ISOFIELDS, parameter)
                    args = (model_size,
                            count_sensors,
                            parameter,
                            pressure_coefficients_storage[angle])
                    self._run_future(executor, self._draw_and_save_plot_future,
                                     PlotBuilding.pseudocolor_coefficients, path, fig_name, 18.5, 10.5, 200, 'tight',
                                     *args)

            executor.shutdown(wait=True)

        # огибающие
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for angle in range(0, angle_border + 5, 5):
                fig_name = f'{model_size_str} {area_type} {angle}'
                path = os.path.join(path_report, ChartType.ENVELOPES)
                args = (pressure_coefficients_storage[angle],
                        (ChartMode.MAX, ChartMode.MEAN, ChartMode.MIN, ChartMode.RMS,
                         ChartMode.STD))
                self._run_future(executor, self._draw_and_save_plot_future,
                                 PlotBuilding.envelopes, path, fig_name, 18.5, 10.5, 200, 'tight',
                                 *args)
            executor.shutdown(wait=True)

        x = np.array(coordinates[0])
        y = np.array(coordinates[1])

        data_to_plot = {}
        data_to_plot_polar = {}

        parameters = [
            ChartMode.MAX,
            ChartMode.MEAN,
            ChartMode.MIN,
            ChartMode.RMS,
            ChartMode.STD,
            ChartMode.CALCULATED,
            ChartMode.WARRANTY_PLUS,
            ChartMode.WARRANTY_MINUS,
        ]

        for v in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
            data_to_plot_polar[v] = {}
            for p in parameters:
                data_to_plot_polar[v][p] = []

        # получаем Cx Cy CMz для всех параметров
        for angle in range(0, angle_border + 5, 5):
            data_to_plot[angle] = {}
            cx, cy = aot_calculations.calculate_cx_cy(
                *count_sensors,
                *size_tpu,
                x,
                y,
                pressure_coefficients_storage[angle]
            )

            cmz = aot_calculations.calculate_cmz(
                *count_sensors,
                angle,
                *size_tpu,
                x,
                y,
                pressure_coefficients_storage[angle]
            )

            for p in parameters:
                data_to_plot_polar[ChartMode.CX][p].append(polar_lambdas[p](cx))
                data_to_plot_polar[ChartMode.CY][p].append(polar_lambdas[p](cy))
                data_to_plot_polar[ChartMode.CMZ][p].append(polar_lambdas[p](cmz))

            data_to_plot[angle][ChartMode.CX] = cx
            data_to_plot[angle][ChartMode.CY] = cy
            data_to_plot[angle][ChartMode.CMZ] = cmz

        # отрисовка CMz Cx Cy
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for angle in range(0, angle_border + 5, 5):
                kt = calculate_kt(model_size, size_tpu, area_type, wind_region, angle)
                fig_name = f'{model_size_str} {area_type} {angle} {ChartMode.CMZ}'

                args = ({ChartMode.CMZ: data_to_plot[angle][ChartMode.CMZ]},
                        kt,
                        self.SAMPLE_PERIOD,
                        self.NUMBER_OF_TIME_COUNTS
                        )
                path = os.path.join(path_report, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS, CoordinateSystem.CARTESIAN)
                self._run_future(executor, self._draw_and_save_plot_future,
                                 PlotBuilding.summary_coefficients, path, fig_name, 18.5, 10.5, 200, 'tight',
                                 *args)

                fig_name = f'{model_size_str} {area_type} {angle} {ChartMode.CX} {ChartMode.CY}'

                args = ({ChartMode.CX: data_to_plot[angle][ChartMode.CX],
                         ChartMode.CY: data_to_plot[angle][ChartMode.CY]},
                        kt,
                        self.SAMPLE_PERIOD,
                        self.NUMBER_OF_TIME_COUNTS
                        )
                self._run_future(executor, self._draw_and_save_plot_future,
                                 PlotBuilding.summary_coefficients, path, fig_name, 18.5, 10.5, 200, 'tight',
                                 *args)

            executor.shutdown(wait=True)

        height = model_size[2]
        speed_sp = speed_sp_region(height, area_type, wind_region)

        # Спектры
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for angle in range(0, angle_border + 5, 5):
                for i in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
                    fig_name = f'{model_size_str} {area_type} {angle}'
                    path = os.path.join(path_report, ChartType.SPECTRUM, i)
                    args = ({i: data_to_plot[angle][i]}, height, speed_sp,
                            self.SAMPLE_FREQUENCY,
                            self.NUMBER_OF_TIME_COUNTS)
                    self._run_future(executor, self._draw_and_save_plot_future,
                                     PlotBuilding.welch_graph, path, fig_name, 18.5, 10.5, 200, 'tight',
                                     *args)
            executor.shutdown(wait=True)

        del data_to_plot

        # Масштабируем данные для полярной системы координат
        if len(pressure_coefficients_storage) != 72:
            for p in parameters:
                cx_scale, cy_scale = scaling_data(data_to_plot_polar[ChartMode.CX][p],
                                                  data_to_plot_polar[ChartMode.CY][p],
                                                  angle_border=angle_border)
                data_to_plot_polar[ChartMode.CX][p] = cx_scale
                data_to_plot_polar[ChartMode.CY][p] = cy_scale

                cmz_scale = scaling_data(data_to_plot_polar[ChartMode.CMZ][p], angle_border=angle_border)
                data_to_plot_polar[ChartMode.CMZ][p] = cmz_scale
        else:
            for p in parameters:
                data_to_plot_polar[ChartMode.CX][p] = np.append(data_to_plot_polar[ChartMode.CX][p],
                                                                data_to_plot_polar[ChartMode.CX][p][0])

                data_to_plot_polar[ChartMode.CY][p] = np.append(data_to_plot_polar[ChartMode.CY][p],
                                                                data_to_plot_polar[ChartMode.CY][p][0])

                data_to_plot_polar[ChartMode.CMZ][p] = np.append(data_to_plot_polar[ChartMode.CMZ][p],
                                                                 data_to_plot_polar[ChartMode.CMZ][p][0])

        # Отрисовка в полярной системе координат
        with ProcessPoolExecutor(max_workers=self.MAX_WORKERS) as executor:
            for i in (ChartMode.CX, ChartMode.CY, ChartMode.CMZ):
                for p in parameters:
                    fig_name = f'{model_size_str} {area_type} {p}'
                    path = os.path.join(path_report, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS,
                                        CoordinateSystem.POLAR, i)
                    args = ({i: {p: data_to_plot_polar[i][p]}},)
                    self._run_future(executor, self._draw_and_save_plot_future,
                                     PlotBuilding.polar_plot, path, fig_name, 18.5, 10.5, 200, 'tight',
                                     *args)
            executor.shutdown(wait=True)

        del data_to_plot_polar

    def create_word_report(
            self,
            model_size,
            wind_region,
            alpha_str,
            path_report,
            report_name

    ):
        breadth, depth, height = model_size
        breadth = int(breadth) if breadth.is_integer() else f'{round(breadth, 2):.2f}'
        depth = int(depth) if depth.is_integer() else f'{round(depth, 2):.2f}'
        height = int(height) if height.is_integer() else f'{round(height, 2):.2f}'

        # Работа с word файлом
        doc = Document()
        style = doc.styles['Normal']
        style.font.size = Pt(14)
        style.font.name = 'Times New Roman'
        section = doc.sections[0]
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        # ширина A4 210 мм высота 297 мм
        fig_width = Mm(165)
        fig_height = Mm(297 / 2 - 55)

        # Шрифт заголовков разного уровня
        head_lvl2 = 16
        head_lvl3 = 16

        counter_plots = 1  # Счетчик графиков для нумерации
        counter_head_lvl1 = 1  # Счетчик заголовков

        WordBuilder.add_heading(doc,
                                head_name=f'Отчет по зданию {breadth}x{depth}x{height}',
                                font_size=24,
                                bold=True)

        for i in ('Параметры ветрового районирования:',
                  f'Ветровой район: {wind_region}',
                  f'Тип местности: {alpha_str}'
                  ):
            doc.add_paragraph().add_run(i)
        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}. Геометрические размеры здания',
                                bold=True)

        counter_head_lvl1 += 1

        table = [
            ['Геометрический размер', 'Значение, м'],
            ['Ширина:', breadth],
            ['Глубина:', depth],
            ['Высота:', height],
        ]
        WordBuilder.add_table(doc, table)

        # Создание содержания
        WordBuilder.add_heading(doc, head_name='Содержание', bold=True, page_break=True)

        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        fldChar = OxmlElement('w:fldChar')  # creates a new element
        fldChar.set(qn('w:fldCharType'), 'begin')  # sets attribute on element
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')  # sets attribute on element
        instrText.text = 'TOC \\o "1-3" \\h \\z \\u'  # change 1-3 depending on heading levels you need

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:t')
        fldChar3.text = "Right-click to update field."
        fldChar2.append(fldChar3)

        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')

        r_element = run._r
        r_element.append(fldChar)
        r_element.append(instrText)
        r_element.append(fldChar2)
        r_element.append(fldChar4)

        # огибающие

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.ENVELOPES}', page_break=True)
        counter_head_lvl1 += 1

        counter_plots = WordBuilder.fill_chapter_with_pictures(
            doc,
            folder_path=os.path.join(path_report, ChartType.ENVELOPES),
            counter_pictures=counter_plots,
            picture_height=fig_height
        )

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.ISOFIELDS}', page_break=True)

        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}. {ChartType.ISOFIELDS} {IsofieldsType.COEFFICIENT.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.COEFFICIENT)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.ISOFIELDS} {IsofieldsType.COEFFICIENT.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)

            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl2 += 1

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.ISOFIELDS} {IsofieldsType.PRESSURE.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.ISOFIELDS, IsofieldsType.PRESSURE)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.ISOFIELDS} {IsofieldsType.PRESSURE.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)
            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl1 += 1
        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS}',
                                page_break=True)

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS} {CoordinateSystem.CARTESIAN.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl2 += 1

        counter_plots = WordBuilder.fill_chapter_with_pictures(
            doc,
            folder_path=os.path.join(path_report, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS,
                                     CoordinateSystem.CARTESIAN),
            counter_pictures=counter_plots,
            picture_height=fig_height - Mm(10)
        )

        WordBuilder.add_heading(doc,
                                head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS} {CoordinateSystem.POLAR.lower()}',
                                head_level=2,
                                font_size=head_lvl2)

        counter_head_lvl3 = 1

        path_temp = os.path.join(path_report, ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS, CoordinateSystem.POLAR)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}.{counter_head_lvl3}. {ChartType.SUMMARY_AERODYNAMIC_COEFFICIENTS} {CoordinateSystem.POLAR.lower()} {mode}',
                                    head_level=3,
                                    font_size=head_lvl3)

            counter_head_lvl3 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        counter_head_lvl1 += 1
        counter_head_lvl2 = 1

        WordBuilder.add_heading(doc, head_name=f'{counter_head_lvl1}. {ChartType.SPECTRUM}', page_break=True)

        path_temp = os.path.join(path_report, ChartType.SPECTRUM)
        for mode in os.listdir(path_temp):
            WordBuilder.add_heading(doc,
                                    head_name=f'{counter_head_lvl1}.{counter_head_lvl2}. {ChartType.SPECTRUM} {mode}',
                                    head_level=2,
                                    font_size=head_lvl2)

            counter_head_lvl2 += 1

            counter_plots = WordBuilder.fill_chapter_with_pictures(
                doc,
                folder_path=os.path.join(path_temp, mode),
                counter_pictures=counter_plots,
                picture_height=fig_height - Mm(10)
            )

        doc.save(os.path.join(path_report, f'{report_name}.docx'))

    def create_report(
            self
    ):
        # need to switch backend to Agg to avoid memory leak
        matplotlib.use('Agg')

        model_size = self._get_model_size()
        alpha_str = self._get_alpha(area_type=True)
        wind_region = self._get_wind_region()

        coordinates = self._get_coordinates()
        face_number = self._get_face_number()

        model_size_str = " ".join(list(map(str, model_size)))
        report_name = f'{model_size_str} {alpha_str} {wind_region}'
        path_report = os.path.join(ReportFolder.WORD_REPORT, self.REPORT_FOLDER_NAME, report_name)

        create_directory_to_report(path_report)
        angle_border = self._get_angle_border()
        size_model_tpu, count_sensors = self._get_size_and_count_sensors(len(coordinates[0]))

        # Получаем коэффициенты сразу для всех углов
        pressure_coefficients_storage = self._get_pressure_coefficients_storage()

        self.create_report_sensor_statistics(coordinates, pressure_coefficients_storage, count_sensors[0],
                                             face_number,
                                             angle_border,
                                             size_model_tpu,
                                             model_size,
                                             path_report)

        self.draw_and_save_all_plots(pressure_coefficients_storage, alpha_str, coordinates, angle_border, path_report,
                                     model_size, model_size_str, size_model_tpu, count_sensors, wind_region)

        del pressure_coefficients_storage

        self.create_word_report(model_size, wind_region, alpha_str, path_report, report_name)
        # return default backend
        matplotlib.use('qtagg')

    def create_report_sensor_statistics(
            self,
            coordinates,
            pressure_coefficients_storage: dict,
            count_sensors,
            face_number,
            angle_border,
            size_model_tpu,
            model_size,
            path_report

    ):
        breadth_tpu, depth_tpu, height_tpu = size_model_tpu
        breadth_real, depth_real, height_real = model_size
        x, z = coordinates

        x_new, y_new = converter_coordinates(x, breadth_tpu, depth_tpu, face_number, count_sensors, accuracy=5)

        headers = (
            'ДАТЧИК',
            'X(м)',
            'Y(м)',
            'Z(м)',
            'Номер грани',
            ChartMode.MEAN,
            ChartMode.RMS,
            ChartMode.STD,
            ChartMode.MAX,
            ChartMode.MIN,
            ChartMode.CALCULATED,
            ChartMode.WARRANTY_PLUS,
            ChartMode.WARRANTY_MINUS
        )
        x_new = [tpu_size_to_real(x_new[sensor], breadth_real, breadth_tpu) for sensor in range(count_sensors)]
        y_new = [tpu_size_to_real(y_new[sensor], depth_real, depth_tpu) for sensor in range(count_sensors)]
        z_real = [tpu_size_to_real(z[sensor], height_real, height_tpu) for sensor in range(count_sensors)]

        wb = Workbook()

        # Удаляем лист по умолчанию, если он не нужен
        default_sheet = wb.active
        wb.remove(default_sheet)

        for angle in range(0, angle_border + 5, 5):
            sheet = wb.create_sheet(f"Угол {angle}")
            sheet.append(headers)
            pressure_coefficients = pressure_coefficients_storage[angle]

            statistics_of_angle = [[i for i in range(1, count_sensors + 1)]]

            statistics_of_angle.append(x_new)
            statistics_of_angle.append(y_new)
            statistics_of_angle.append(z_real)
            statistics_of_angle.append(face_number)

            for k in lambdas:
                statistics_of_angle.append(lambdas[k](pressure_coefficients))

            statistics_of_angle.append(calculated(pressure_coefficients, axis=0))
            statistics_of_angle.append(warranty_plus(pressure_coefficients, axis=0))
            statistics_of_angle.append(warranty_minus(pressure_coefficients, axis=0))

            statistics_of_angle = np.array(statistics_of_angle).T

            for row in statistics_of_angle:
                sheet.append(row.tolist())

        wb.save(f'{os.path.join(path_report, ReportFolder.FILE_NAME_SENSOR_STATISTICS)}.xlsx')
