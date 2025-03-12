import os
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pydantic import validate_call


class WordBuilder:
    @staticmethod
    @validate_call
    def add_picture(
            doc,
            picture_path: str,
            picture_height: int | None = None,
            picture_width: int | None = None,
            chapter: str = '',
            picture_name: str = ''
    ):
        p = doc.add_paragraph()
        run = p.add_run()
        run.add_picture(picture_path, height=picture_height, width=picture_width)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = doc.add_paragraph(f'Рисунок {chapter} {picture_name}')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    @validate_call
    def add_heading(
            doc,
            head_name: str = '',
            head_level: int = 1,
            font_size: int = 20,
            bold: bool = False,
            page_break: bool = False
    ):
        if page_break:
            doc.add_page_break()

        head = doc.add_heading(level=head_level)
        run = head.add_run(head_name)
        run.font.size = Pt(font_size)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.bold = bold

    @staticmethod
    @validate_call
    def add_table(
            doc,
            table
    ):
        count_cols = len(table[0])
        doc_table = doc.add_table(rows=len(table), cols=count_cols)
        doc_table.style = 'Table Grid'

        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                doc_table.cell(i, j).text = str(cell)

        for i in range(count_cols):
            doc_table.cell(0, i).paragraphs[0].runs[0].bold = True

    @staticmethod
    @validate_call
    def fill_chapter_with_pictures(
            doc,
            folder_path: str,
            counter_pictures: int = 1,
            picture_height: int | None = None,
            picture_width: int | None = None,
    ) -> int:
        counter_page_break = 0

        for plot in os.listdir(folder_path):
            WordBuilder.add_picture(
                doc,
                picture_path=os.path.join(folder_path, plot),
                picture_height=picture_height,
                picture_width=picture_width,
                chapter=f'{counter_pictures}.',
                picture_name=Path(plot).stem
            )

            counter_pictures += 1
            counter_page_break += 1

            if counter_page_break % 2 == 0:
                doc.add_page_break()

        return counter_pictures
